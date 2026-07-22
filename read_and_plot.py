"""
Penetrometer Data Logger
------------------------
1  → start recording
2  → stop recording, show plot, save Excel

Requirements:
    pip install pyserial openpyxl matplotlib
"""

import serial
import serial.tools.list_ports
import time
import os
from datetime import datetime

import matplotlib.pyplot as plt
import openpyxl
from openpyxl.chart import LineChart, Reference

# ── Config ────────────────────────────────────────────────────────
PORT     = None    # None = auto-detect; or set e.g. "COM4"
BAUD     = 115200
SAVE_DIR = os.path.dirname(os.path.abspath(__file__))

# ── Auto-detect port ──────────────────────────────────────────────
def find_port():
    ports = serial.tools.list_ports.comports()
    for p in ports:
        if any(kw in (p.description or "").upper()
               for kw in ["ESP32", "CP210", "CH340", "USB SERIAL", "FTDI"]):
            return p.device
    return ports[0].device if ports else None

# ── Save Excel ────────────────────────────────────────────────────
def save_excel(times, forces, qus_kpa, qus_kgcm2):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Penetrometer Data"

    ws.append(["Time (s)", "Force (kg)", "qu (kPa)", "qu (kg/cm2)"])
    for row in zip(times, forces, qus_kpa, qus_kgcm2):
        ws.append(list(row))

    n = len(times) + 1

    # Force chart
    c1 = LineChart()
    c1.title = "Force vs Time"
    c1.y_axis.title = "Force (kg)"
    c1.x_axis.title = "Time (s)"
    c1.width = 20; c1.height = 12
    c1.add_data(Reference(ws, min_col=2, min_row=1, max_row=n), titles_from_data=True)
    c1.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
    ws.add_chart(c1, "F2")

    # qu chart
    c2 = LineChart()
    c2.title = "qu vs Time"
    c2.y_axis.title = "qu (kPa)"
    c2.x_axis.title = "Time (s)"
    c2.width = 20; c2.height = 12
    c2.add_data(Reference(ws, min_col=3, min_row=1, max_row=n), titles_from_data=True)
    c2.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
    ws.add_chart(c2, "F22")

    fname = os.path.join(SAVE_DIR,
        f"penetrometer_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    wb.save(fname)
    print(f"Excel saved → {fname}")

# ── Plot results ──────────────────────────────────────────────────
def show_plot(times, forces, qus_kpa):
    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(10, 7), sharex=True)
    fig.suptitle("Penetrometer Results", fontsize=13)

    ax1.plot(times, forces, color="steelblue", linewidth=1.5)
    ax1.set_ylabel("Force (kg)")
    ax1.grid(True)

    ax2.plot(times, qus_kpa, color="darkorange", linewidth=1.5)
    ax2.set_ylabel("qu (kPa)")
    ax2.set_xlabel("Time (s)")
    ax2.grid(True)

    plt.tight_layout()
    plt.show()

# ── Main ──────────────────────────────────────────────────────────
def main():
    port = PORT or find_port()
    if not port:
        print("No serial port found. Set PORT manually in the script.")
        return

    ser = serial.Serial(port, BAUD, timeout=2)
    print(f"Connected to {port}\n")

    # ── Pass-through: calibration wizard ─────────────────────────
    print("Complete calibration below (your typing is sent to the ESP32).")
    print("When you see readings appearing, press 1 to START recording.\n")

    import threading

    header_found = threading.Event()

    def forward_input():
        """Forward keyboard input to ESP32 during calibration."""
        while not header_found.is_set():
            try:
                user_input = input()
                if not header_found.is_set():
                    ser.write((user_input + "\n").encode("utf-8"))
            except EOFError:
                break

    input_thread = threading.Thread(target=forward_input, daemon=True)
    input_thread.start()

    while not header_found.is_set():
        line = ser.readline().decode("utf-8", errors="ignore").strip()
        if line:
            print(line)
        if line.startswith("TIME_MS"):
            header_found.set()

    input_thread.join(timeout=0.5)

    # ── Wait for user to press 1 ──────────────────────────────────
    print("\n>>> Press 1 + Enter to START recording  |  2 + Enter to STOP\n")
    while True:
        cmd = input().strip()
        if cmd == "1":
            break
        print("Type 1 to start.")

    # ── Record ────────────────────────────────────────────────────
    print("Recording... press 2 + Enter to stop.\n")
    times     = []
    forces    = []
    qus_kpa   = []
    qus_kgcm2 = []

    stop_flag = threading.Event()

    def wait_for_stop():
        while True:
            cmd = input().strip()
            if cmd == "2":
                stop_flag.set()
                break

    t = threading.Thread(target=wait_for_stop, daemon=True)
    t.start()

    while not stop_flag.is_set():
        line = ser.readline().decode("utf-8", errors="ignore").strip()
        if not line or line.startswith("TIME_MS"):
            continue
        parts = line.split(",")
        if len(parts) != 4:
            continue
        try:
            times.append(float(parts[0]) / 1000.0)
            forces.append(float(parts[1]))
            qus_kpa.append(float(parts[2]))
            qus_kgcm2.append(float(parts[3]))
        except ValueError:
            continue

    ser.close()
    print(f"\nRecording stopped. {len(times)} samples collected.")

    if not times:
        print("No data collected.")
        return

    # Normalize time to start at 0
    t0 = times[0]
    times = [t - t0 for t in times]

    save_excel(times, forces, qus_kpa, qus_kgcm2)
    show_plot(times, forces, qus_kpa)

if __name__ == "__main__":
    main()
