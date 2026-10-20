"""
Record force vs time for the 1.5 inch tip. Requires calibration.txt.

  1 + Enter -> zero (1 s, hands off the probe) and start recording
  2 + Enter -> stop -> CSV + Excel + quick-look plot  (saved in data/)

Differences from the original record.py (same workflow otherwise):
  - time comes from the ESP32 clock, so sample spacing is accurate
  - the zero baseline is taken right when you press 1 (1 s average)
  - force is kept signed (no abs), so noise is not biased upward; if the
    load cell reads negative under load the whole record is flipped once
  - CSV also stores Force (N), Pressure (kPa), UCS (kPa) +/- uncertainty
    and the Raw ADC value (UCS = tip pressure, pocket-penetrometer method)
  - a <name>_meta.json file records calibration, baseline, warnings
  - live readout and warnings for ADC saturation / stuck signal
"""

import csv, json, os, threading, time
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.chart import ScatterChart, Reference, Series

from common import (DATA_DIR, TIP_AREA_M2, TIP_DIAMETER_IN, ensure_dir, timestamp, safe_label,
                    kg_to_N, kg_to_kPa, ucs_uncertainty_kPa, load_calibration, load_characterization, calibration_age_days,
                    parse_args, open_serial, Streamer, raw_health, max_measurable_kg, open_file,
                    ADC_MAX, ADC_MIN, SAT_FRAC)

BASELINE_S = 1.0

args = parse_args("Record a 1.5\" tip penetrometer test")

ZERO_OFFSET, CAL_FACTOR = load_calibration()
char = load_characterization()
print(f"Loaded calibration: zero={ZERO_OFFSET}, factor={CAL_FACTOR:.9g} kg/count")
age = calibration_age_days()
if age is not None and age > 1:
    print(f"  NOTE: calibration.txt is {age:.0f} days old.")
if char:
    print(f"System characterization ({char['date']}): max error {char['max_abs_error_kg']*1000:.0f} g, "
          f"noise {char['noise_sigma_kg']*1000:.1f} g, usable up to {char['usable_max_kg']:.0f} kg "
          f"({char['usable_max_kPa']:.0f} kPa)")
print(f"Tip: {TIP_DIAMETER_IN}\" diameter, area {TIP_AREA_M2*1e6:.1f} mm^2  (1 kg = {kg_to_kPa(1):.2f} kPa)\n")

label = safe_label(input("Sample label (e.g. wet_grey_s1, dry_red_s2): "))

ser = open_serial(args.port)
st  = Streamer(ser)
st.start()
print("Streaming OK.")

print("\nType  1 + Enter  to ZERO and START   (keep hands off the probe while it zeros)")
print("Type  2 + Enter  to STOP\n")
while input().strip() != "1":
    print("Type 1 to start.")

# ── Baseline (zero) ───────────────────────────────────────────────────
base = st.collect(BASELINE_S, "zeroing - hands off")
base_raw = np.array([r for _, r in base], float)
baseline_kg = float(np.mean((base_raw - ZERO_OFFSET) * CAL_FACTOR))
baseline_sd = float(np.std((base_raw - ZERO_OFFSET) * CAL_FACTOR))
warnings = []
noise_ref = char["noise_sigma_kg"] if char else None
if noise_ref and baseline_sd > 4 * noise_ref:
    warnings.append(f"baseline noisy ({baseline_sd*1000:.1f} g vs {noise_ref*1000:.1f} g normal) - probe touched while zeroing?")
if abs(baseline_kg) > 1.0:
    warnings.append(f"zero has shifted {baseline_kg:+.2f} kg since calibration (load on tip at start? temperature?)")
for w in warnings:
    print("WARNING:", w)

# Highest load this run can measure before the ADC saturates
headroom_kg = max_measurable_kg(float(base_raw.mean()), CAL_FACTOR)

start_index = st.count()
print(f"\nRECORDING  (zero = {baseline_kg:+.4f} kg)   - type 2 + Enter to stop\n")

stop = threading.Event()


def wait_stop():
    while input().strip() != "2":
        pass
    stop.set()


threading.Thread(target=wait_stop, daemon=True).start()

peak = 0.0
warned = set()
last_print = 0.0
while not stop.is_set():
    time.sleep(0.05)
    if st.error:
        print(f"\nERROR: connection lost ({st.error}). Saving what was recorded.")
        break
    if time.time() - last_print >= 1.0:
        last_print = time.time()
        recent = st.since(max(start_index, st.count() - 40))
        if not recent:
            continue
        raws = [r for _, r in recent]
        f_now = (np.mean(raws[-4:]) - ZERO_OFFSET) * CAL_FACTOR - baseline_kg
        peak  = max(peak, abs(f_now))
        t_now = (recent[-1][0] - base[-1][0]) / 1000.0 if base else 0
        print(f"  t={t_now:6.1f} s   F={f_now:8.3f} kg  {kg_to_N(f_now):8.2f} N  "
              f"UCS {kg_to_kPa(f_now):8.2f} kPa   peak UCS {kg_to_kPa(peak):8.2f} kPa")
        for w in raw_health(raws):
            if w not in warned:
                print("  WARNING:", w); warned.add(w)
        if abs(f_now) > 0.9 * abs(headroom_kg) and "near_limit" not in warned:
            print("  WARNING: above 90 % of the measurable range!"); warned.add("near_limit")

st.stop()
ser.close()

samples = st.since(start_index)
print(f"\nStopped. {len(samples)} samples.\n")
if not samples:
    print("No data collected.")
    raise SystemExit(0)

# ── Convert ───────────────────────────────────────────────────────────
ms  = np.array([s[0] for s in samples], float)
raw = np.array([s[1] for s in samples], float)
times  = (ms - ms[0]) / 1000.0
force  = (raw - ZERO_OFFSET) * CAL_FACTOR - baseline_kg

flipped = False
if abs(force.min()) > abs(force.max()):        # load cell reads negative under load
    force = -force
    flipped = True
    print("Note: force was negative under load - sign flipped so load is positive.")

back = np.diff(ms)
if len(back) and back.min() < 0:
    warnings.append("timestamps went backwards - the board probably reset during the test")
if len(back) and np.median(back) > 0 and back.max() > 10 * np.median(back):
    warnings.append(f"data gap of {back.max()/1000:.2f} s in the record")
warnings += raw_health(raw.astype(int).tolist())

forces_N   = kg_to_N(force)
pressure   = kg_to_kPa(force)
ucs        = pressure                                  # UCS estimate = pressure on the 1.5" tip
ucs_u      = ucs_uncertainty_kPa(force, char)          # +/- per sample (None without characterization)
ucs_u_col  = ucs_u if ucs_u is not None else np.full_like(ucs, np.nan)
rate_hz    = (len(times) - 1) / times[-1] if times[-1] > 0 else 0

ts        = timestamp()
ensure_dir(DATA_DIR)
base_name = os.path.join(DATA_DIR, f"{label}_{ts}")
csv_name  = base_name + ".csv"
xlsx_name = base_name + ".xlsx"
png_name  = base_name + ".png"
meta_name = base_name + "_meta.json"

# ── CSV (first two columns identical to the original format) ─────────
COLUMNS = ["Time (s)", "Force (kg)", "Force (N)", "Pressure (kPa)",
           "UCS (kPa)", "UCS uncertainty (+/- kPa)", "Raw ADC"]
with open(csv_name, "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(COLUMNS)
    for row in zip(times, force, forces_N, pressure, ucs, ucs_u_col, raw):
        w.writerow([f"{row[0]:.3f}", f"{row[1]:.5f}", f"{row[2]:.4f}", f"{row[3]:.4f}",
                    f"{row[4]:.4f}", "" if np.isnan(row[5]) else f"{row[5]:.4f}", int(row[6])])
print(f"CSV saved   -> {csv_name}")

meta = {
    "label": label, "timestamp": ts, "samples": int(len(times)), "duration_s": float(times[-1]),
    "sample_rate_hz": float(rate_hz),
    "tip_diameter_in": TIP_DIAMETER_IN, "tip_area_m2": TIP_AREA_M2,
    "calibration_zero": ZERO_OFFSET, "calibration_factor": CAL_FACTOR,
    "baseline_kg": baseline_kg, "baseline_sd_kg": baseline_sd, "sign_flipped": flipped,
    "peak_kg": float(force.max()), "peak_N": float(forces_N.max()), "peak_kPa": float(pressure.max()),
    "peak_UCS_kPa": float(ucs.max()),
    "peak_UCS_uncertainty_kPa": float(ucs_u[int(np.argmax(ucs))]) if ucs_u is not None else None,
    "measurable_max_kg": float(abs(headroom_kg)),
    "warnings": warnings,
    "characterization_date": char["date"] if char else None,
}
with open(meta_name, "w") as f:
    json.dump(meta, f, indent=2)

# ── Excel ─────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Data"
ws.append(COLUMNS)
for row in zip(times, force, forces_N, pressure, ucs, ucs_u_col, raw):
    ws.append([round(float(row[0]), 3), float(row[1]), float(row[2]), float(row[3]), float(row[4]),
               None if np.isnan(row[5]) else float(row[5]), int(row[6])])
n = len(times) + 1
for col, title, anchor in [(3, "Force (N)", "I2"), (5, "UCS (kPa)", "I28")]:
    ch = ScatterChart()
    ch.title = f"{title} vs Time"
    ch.x_axis.title = "Time (s)"; ch.y_axis.title = title
    ch.width = 22; ch.height = 12
    s = Series(Reference(ws, min_col=col, min_row=1, max_row=n),
               Reference(ws, min_col=1, min_row=2, max_row=n), title_from_data=True)
    s.marker.symbol = "none"; s.smooth = False
    ch.series.append(s)
    ws.add_chart(ch, anchor)
info = wb.create_sheet("Info")
for k, v in meta.items():
    info.append([k, json.dumps(v) if isinstance(v, (list, dict)) else v])
wb.save(xlsx_name)
print(f"Excel saved -> {xlsx_name}")

# ── Quick-look plot ──────────────────────────────────────────────────
win = max(1, int(round(rate_hz * 0.25)))                     # 0.25 s moving average
kernel = np.ones(win) / win
f_s = np.convolve(forces_N, kernel, mode="same") if len(forces_N) > win else forces_N
dfdt = np.gradient(f_s, times) if len(times) > 2 else np.zeros_like(times)

fig, (ax1, axu, ax2) = plt.subplots(3, 1, figsize=(11, 10.5), sharex=True,
                                    gridspec_kw={"height_ratios": [3, 3, 2]})
ax1.plot(times, forces_N, color="steelblue", lw=1)
i_pk = int(np.argmax(forces_N))
ax1.plot(times[i_pk], forces_N[i_pk], "ro")
ax1.annotate(f"peak {forces_N[i_pk]:.1f} N = {pressure[i_pk]:.1f} kPa", (times[i_pk], forces_N[i_pk]),
             textcoords="offset points", xytext=(8, -12))
ax1.set_ylabel("Force (N)")
ax1.set_title(f"{label} - 1.5\" tip  -  force vs time")
ax1.grid(True)
if ucs_u is not None:
    axu.fill_between(times, ucs - ucs_u, ucs + ucs_u, color="seagreen", alpha=0.2, lw=0,
                     label="+/- calibration uncertainty")
axu.plot(times, ucs, color="seagreen", lw=1, label="UCS estimate")
axu.plot(times[i_pk], ucs[i_pk], "ro")
u_txt = f" +/- {ucs_u[i_pk]:.1f}" if ucs_u is not None else ""
axu.annotate(f"peak UCS {ucs[i_pk]:.1f}{u_txt} kPa", (times[i_pk], ucs[i_pk]),
             textcoords="offset points", xytext=(8, -12))
axu.set_ylabel("UCS (kPa)")
axu.set_title("Unconfined compressive strength estimate vs time")
axu.grid(True); axu.legend(loc="upper left")
ax2.plot(times, dfdt, color="darkorange", lw=1)
ax2.axhline(0, color="black", lw=0.8, ls="--")
ax2.set_ylabel("dF/dt (N/s, smoothed)")
ax2.set_xlabel("Time (s)")
ax2.grid(True)
if warnings:
    fig.text(0.01, 0.005, "WARNINGS: " + " | ".join(warnings), color="red", fontsize=8)
fig.tight_layout(rect=(0, 0.02, 1, 1))
fig.savefig(png_name, dpi=150)
print(f"Plot saved  -> {png_name}")

print(f"\nPeak: {forces_N.max():.2f} N   UCS = {ucs.max():.2f}{u_txt} kPa   "
      f"({len(times)} samples, {times[-1]:.1f} s, {rate_hz:.0f} Hz)")
if warnings:
    print("\n!!! WARNINGS for this test:")
    for w in warnings:
        print("   -", w)
open_file(png_name)
